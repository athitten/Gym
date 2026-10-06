# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""A configured ``libreoffice_command`` must drive every Office -> PDF conversion.

The fake converter is a real executable, so these tests exercise argv
construction, executable lookup, and output discovery the way production does.
PATH is emptied where it matters, so a host ``libreoffice`` cannot satisfy them.
"""

import base64
import shutil
import stat
from pathlib import Path

import pytest

from resources_servers.gdpval import preconvert, setup_libreoffice
from responses_api_agents.stirrup_agent import file_reader


_SCRIPT = """#!/bin/sh
PATH=/usr/bin:/bin
printf '%s\\n' "$*" >> '@LOG@'
if [ "$1" = "--version" ]; then echo "LibreOffice 99.9 fake"; exit @VERSION_RC@; fi
outdir=""; src=""
while [ $# -gt 0 ]; do
  case "$1" in --outdir) outdir="$2"; shift ;; *) src="$1" ;; esac
  shift
done
base=$(basename "$src")
@CONVERT@
"""


@pytest.fixture
def template_pdf(tmp_path: Path) -> Path:
    import fitz

    path = tmp_path / "template.pdf"
    document = fitz.open()
    document.new_page().insert_text((72, 72), "rendered by the configured converter")
    document.save(path)
    document.close()
    return path


def _converter(directory: Path, template: Path, name: str = "fake-lo", *, converts=True, version_rc=0):
    """Write an executable that logs argv and writes ``<outdir>/<stem>.pdf`` like LibreOffice."""
    directory.mkdir(parents=True, exist_ok=True)
    log = directory / f"{name}.log"
    convert = f"cp '{template}' \"$outdir/${{base%.*}}.pdf\"" if converts else ":"
    script = directory / name
    script.write_text(
        _SCRIPT.replace("@LOG@", str(log)).replace("@VERSION_RC@", str(version_rc)).replace("@CONVERT@", convert)
    )
    script.chmod(script.stat().st_mode | stat.S_IXUSR)
    return script, log


def _calls(log: Path) -> list[str]:
    return log.read_text().splitlines() if log.exists() else []


@pytest.fixture
def no_host_libreoffice(tmp_path: Path, monkeypatch) -> None:
    empty = tmp_path / "empty-path"
    empty.mkdir()
    monkeypatch.setenv("PATH", str(empty))


def test_ensure_accepts_working_command_without_host_checks(tmp_path, template_pdf, monkeypatch):
    script, log = _converter(tmp_path / "bin", template_pdf)
    real_run = setup_libreoffice._run
    commands: list[list[str]] = []

    def _spy(cmd, **kwargs):
        commands.append(cmd)
        return real_run(cmd, **kwargs)

    def _host_check(*_args):
        raise AssertionError("the host javaldx/java check ran for a configured command")

    monkeypatch.setattr(setup_libreoffice, "_run", _spy)
    monkeypatch.setattr(setup_libreoffice, "_javaldx_works", _host_check)
    monkeypatch.setattr(setup_libreoffice, "_java_runs", _host_check)

    assert setup_libreoffice.ensure_libreoffice([str(script)]) is True

    version, conversion = _calls(log)
    assert version == "--version"
    assert "--convert-to pdf" in conversion and conversion.endswith("probe.xlsx")
    assert not any("apt-get" in cmd[0] for cmd in commands)


@pytest.mark.parametrize("converts, version_rc", [(False, 0), (True, 1)])
def test_ensure_rejects_command_that_fails(tmp_path, template_pdf, converts, version_rc):
    script, _ = _converter(tmp_path / "bin", template_pdf, converts=converts, version_rc=version_rc)

    assert setup_libreoffice.ensure_libreoffice([str(script)]) is False


def test_ensure_rejects_command_whose_output_is_not_a_pdf(tmp_path):
    # A converter can exit 0 and leave a .pdf that is really an error page; the probe checks the bytes.
    not_pdf = tmp_path / "error-page.html"
    not_pdf.write_text("<html>conversion failed</html>")
    script, log = _converter(tmp_path / "bin", not_pdf)

    assert setup_libreoffice.ensure_libreoffice([str(script)]) is False
    assert len(_calls(log)) == 2


def test_ensure_rejects_missing_command(tmp_path):
    assert setup_libreoffice.ensure_libreoffice([str(tmp_path / "missing")]) is False


def test_preconvert_dir_uses_command_on_every_path(tmp_path, template_pdf, no_host_libreoffice):
    script, log = _converter(tmp_path / "bin", template_pdf)
    deliverables = tmp_path / "deliverables"
    deliverables.mkdir()
    # Parallel path, whitespace staging, and the serial same-stem sidecar path.
    for name in ("memo.docx", "Q1 notes.docx", "Plan.docx", "Plan.pptx"):
        (deliverables / name).write_bytes(b"office source")

    result = preconvert.preconvert_dir(deliverables, libreoffice_command=[str(script)])

    assert result == (4, 0, [])
    for name in ("memo.pdf", "Q1 notes.pdf", "Plan.docx.pdf", "Plan.pptx.pdf"):
        assert (deliverables / name).read_bytes() == template_pdf.read_bytes()
    assert len(_calls(log)) == 4


def test_rubric_rendering_uses_command_instead_of_text_fallback(tmp_path, template_pdf, no_host_libreoffice):
    from docx import Document

    script, _ = _converter(tmp_path / "bin", template_pdf)
    deliverables = tmp_path / "deliverables"
    deliverables.mkdir()
    document = Document()
    document.add_paragraph("Quarterly throughput rose 12 percent.")
    document.save(deliverables / "report.docx")

    configured = file_reader.convert_deliverables_to_content_blocks(
        str(deliverables), libreoffice_command=[str(script)]
    )
    default = file_reader.convert_deliverables_to_content_blocks(str(deliverables))

    texts = [block["text"] for block in configured if block["type"] == "text"]
    assert "\nreport.docx (converted to PDF):" in texts
    pdfs = [block["image_url"]["url"] for block in configured if block["type"] == "image_url"]
    assert pdfs == ["data:application/pdf;base64," + base64.b64encode(template_pdf.read_bytes()).decode()]
    # Rendering happens in scratch space; the deliverables directory is untouched.
    assert sorted(p.name for p in deliverables.iterdir()) == ["report.docx"]
    # Without the command the same host degrades to text, as before.
    assert any("(text fallback)" in block.get("text", "") for block in default)


def test_unconfigured_default_runs_host_libreoffice_with_historical_argv(tmp_path, template_pdf, monkeypatch):
    bin_dir = tmp_path / "bin"
    _, log = _converter(bin_dir, template_pdf, name="libreoffice")
    monkeypatch.setenv("PATH", str(bin_dir))
    source = tmp_path / "work" / "memo.docx"
    source.parent.mkdir()
    source.write_bytes(b"office source")

    _, ok, _ = preconvert.convert_to_pdf(source)
    rendered = file_reader._convert_office_to_pdf(source, out_dir=tmp_path)

    assert ok and rendered == tmp_path / "memo.pdf"
    for call, outdir in zip(_calls(log), (source.parent, tmp_path), strict=True):
        parts = call.split(" ")
        assert parts[:5] == ["--headless", "--nologo", "--nolockcheck", "--nodefault", "--norestore"]
        assert parts[5].startswith("-env:UserInstallation=file:///")
        assert parts[6:] == ["--convert-to", "pdf", "--outdir", str(outdir), str(source)]


@pytest.mark.skipif(shutil.which("libreoffice") is None, reason="LibreOffice is not installed")
def test_real_libreoffice_through_wrapper_command(tmp_path):
    from docx import Document

    wrapper = tmp_path / "lo-wrapper"
    wrapper.write_text('#!/bin/sh\nexec libreoffice "$@"\n')
    wrapper.chmod(wrapper.stat().st_mode | stat.S_IXUSR)
    deliverables = tmp_path / "deliverables"
    deliverables.mkdir()
    document = Document()
    document.add_paragraph("Rendered through a wrapper command.")
    document.save(deliverables / "report.docx")

    assert setup_libreoffice.ensure_libreoffice([str(wrapper)]) is True
    assert preconvert.preconvert_dir(deliverables, libreoffice_command=[str(wrapper)]) == (1, 0, [])
    assert (deliverables / "report.pdf").read_bytes().startswith(b"%PDF")
